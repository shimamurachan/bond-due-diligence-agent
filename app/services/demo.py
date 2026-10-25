from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from docx import Document as DocxDocument


@dataclass(frozen=True)
class DemoMaterial:
    file_name: str
    source_note: str
    content: bytes


def build_demo_materials() -> list[DemoMaterial]:
    """Return intentionally fictional DOCX material for an isolated trial project."""
    return [
        DemoMaterial(
            "远景实业集团2026年度第一期公司债券募集说明书（脱敏试用）.docx",
            "系统生成的脱敏试用材料，不代表真实发行人。",
            _docx_bytes(
                "远景实业集团 2026 年度第一期公司债券募集说明书（节选）",
                [
                    "发行人：远景实业集团有限公司。本期拟发行公司债券，期限 3 年，发行规模不超过人民币 10 亿元。",
                    "募集资金用途：其中 6 亿元用于偿还到期有息负债，3 亿元用于补充营运资金，1 亿元用于项目建设。募集资金应按约定用途专户管理。",
                    "偿债资金来源包括经营活动现金流、货币资金及必要时的融资安排。",
                ],
                [("项目", "金额（亿元）"), ("偿还到期有息负债", "6"), ("补充营运资金", "3"), ("项目建设", "1")],
            ),
        ),
        DemoMaterial(
            "远景实业集团2025年度财务数据摘要（脱敏试用）.docx",
            "系统生成的脱敏试用材料，不代表真实财务数据。",
            _docx_bytes(
                "远景实业集团 2025 年度财务数据摘要（节选）",
                [
                    "截至 2025 年末，发行人合并口径资产总额 120 亿元，负债总额 92 亿元，资产负债率 76.7%。",
                    "2025 年经营活动产生的现金流量净额为 -1.8 亿元；期末货币资金为 5.2 亿元，其中受限资金 2.4 亿元。",
                    "发行人未来十二个月到期有息负债为 18 亿元。本摘要未提供授信额度及续作融资安排明细。",
                ],
                [("指标", "2025 年末或全年"), ("资产负债率", "76.7%"), ("经营活动现金流净额", "-1.8 亿元"), ("未来十二个月到期有息负债", "18 亿元")],
            ),
        ),
        DemoMaterial(
            "担保增信说明（脱敏试用）.docx",
            "系统生成的脱敏试用材料，不代表真实担保文件。",
            _docx_bytes(
                "担保增信说明（节选）",
                [
                    "本期债券拟由远景控股有限公司提供连带责任保证担保。",
                    "材料仅列示担保人名称，未提供担保人最近一期财务报表、对外担保余额、保证期间及保证范围的完整条款。",
                    "发行人应在发行前补充担保合同及担保人偿债能力说明。",
                ],
                [],
            ),
        ),
    ]


def _docx_bytes(title: str, paragraphs: list[str], rows: list[tuple[str, str]]) -> bytes:
    document = DocxDocument()
    document.add_heading(title, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if rows:
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for left, right in rows:
            cells = table.add_row().cells
            cells[0].text = left
            cells[1].text = right
    output = BytesIO()
    document.save(output)
    return output.getvalue()
