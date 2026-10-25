from __future__ import annotations

import hashlib
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

import fitz
import jieba
from docx import Document as DocxDocument

from app.config import Settings, get_settings


@dataclass(frozen=True)
class ParsedPage:
    page_number: int
    content: str
    extraction_method: str  # native | ocr | docx
    ocr_confidence: float | None = None


@dataclass(frozen=True)
class ParsedChunk:
    page_start: int
    page_end: int
    location: str
    structure_type: str
    content: str
    search_terms: str
    content_hash: str


class DocumentParseError(RuntimeError):
    pass


def clean_text(value: str) -> str:
    return "\n".join(line.strip() for line in value.replace("\u00a0", " ").splitlines() if line.strip()).strip()


def tokenise_for_search(value: str) -> str:
    return " ".join(token for token in jieba.lcut(value) if token.strip())


class DocumentParser:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()

    def parse(self, file_name: str, content_type: str, data: bytes) -> list[ParsedPage]:
        lowered = file_name.lower()
        if lowered.endswith(".pdf") or content_type == "application/pdf":
            return self._parse_pdf(data)
        if lowered.endswith(".docx") or "wordprocessingml" in content_type:
            return self._parse_docx(data)
        raise DocumentParseError("目前仅支持 PDF 与 DOCX 文件。")

    def _parse_pdf(self, data: bytes) -> list[ParsedPage]:
        try:
            pdf = fitz.open(stream=data, filetype="pdf")
        except Exception as exc:
            raise DocumentParseError("PDF 无法打开，可能已损坏或受密码保护。") from exc
        if len(pdf) == 0:
            raise DocumentParseError("PDF 不含可处理页面。")
        if len(pdf) > self.settings.max_pdf_pages:
            raise DocumentParseError(f"PDF 超过 {self.settings.max_pdf_pages} 页上限。")
        try:
            pages: list[ParsedPage] = []
            with tempfile.TemporaryDirectory(prefix="bond-ocr-") as temp_dir:
                workspace = Path(temp_dir)
                for number, page in enumerate(pdf, start=1):
                    native_text = clean_text(page.get_text("text"))
                    if len(native_text) >= 80:
                        pages.append(ParsedPage(number, native_text, "native"))
                        continue
                    image_path = workspace / f"page-{number}.png"
                    scale = self.settings.ocr_dpi / 72
                    page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).save(image_path)
                    ocr_text, confidence = self._ocr_page(image_path, workspace / f"ocr-{number}")
                    if not ocr_text:
                        raise DocumentParseError(f"第 {number} 页无法从扫描件中识别出文字。")
                    pages.append(ParsedPage(number, clean_text(ocr_text), "ocr", confidence))
            return pages
        finally:
            pdf.close()

    def _parse_docx(self, data: bytes) -> list[ParsedPage]:
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temporary:
            temporary.write(data)
            path = Path(temporary.name)
        try:
            document = DocxDocument(path)
            blocks: list[str] = []
            for paragraph in document.paragraphs:
                value = clean_text(paragraph.text)
                if value:
                    blocks.append(value)
            for table in document.tables:
                rows = []
                for row in table.rows:
                    cells = [clean_text(cell.text).replace("|", "\\|") for cell in row.cells]
                    if any(cells):
                        rows.append("| " + " | ".join(cells) + " |")
                if rows:
                    if len(rows) > 1:
                        header_count = rows[0].count("|") - 1
                        rows.insert(1, "| " + " | ".join(["---"] * header_count) + " |")
                    blocks.append("\n".join(rows))
            content = "\n\n".join(blocks)
            if not content:
                raise DocumentParseError("DOCX 未包含可检索正文或表格。")
            return [ParsedPage(1, content, "docx")]
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError("DOCX 解析失败。") from exc
        finally:
            path.unlink(missing_ok=True)

    def _ocr_page(self, image_path: Path, output_dir: Path) -> tuple[str, float | None]:
        """Use PaddleOCR PP-StructureV3 so scan pages preserve tables/reading order as Markdown."""
        output_dir.mkdir(parents=True, exist_ok=True)
        runner = Path(__file__).resolve().parents[2] / "scripts" / "run_ocr.py"
        command = [
            sys.executable, str(runner), str(image_path), str(output_dir),
        ]
        try:
            completed = subprocess.run(command, capture_output=True, text=True, timeout=180, check=False)
        except FileNotFoundError as exc:
            raise DocumentParseError("PaddleOCR 未安装，无法处理扫描件。") from exc
        except subprocess.TimeoutExpired as exc:
            raise DocumentParseError("扫描页 OCR 超时，请拆分文件后重试。") from exc
        if completed.returncode != 0:
            raise DocumentParseError(f"扫描页 OCR 失败：{completed.stderr[-300:] or completed.stdout[-300:]}")
        markdown_files = sorted(output_dir.rglob("*.md"))
        if not markdown_files:
            raise DocumentParseError("OCR 未生成可用的 Markdown 结果。")
        text = "\n\n".join(path.read_text(encoding="utf-8", errors="ignore") for path in markdown_files)
        # PP-Structure does not guarantee one global score in CLI output. The stored method is
        # still enough for the UI to disclose OCR-derived evidence; low-confidence per-line data
        # stays in the raw Paddle result rather than being fabricated here.
        return text, None


def build_chunks(pages: list[ParsedPage], target_size: int = 800, overlap: int = 120) -> list[ParsedChunk]:
    """Page-aware chunks; Markdown tables are kept intact wherever possible."""
    chunks: list[ParsedChunk] = []
    for page in pages:
        paragraphs = [part.strip() for part in page.content.split("\n\n") if part.strip()]
        current = ""
        part = 1
        for paragraph in paragraphs:
            is_table = paragraph.lstrip().startswith("|")
            if is_table:
                if current:
                    chunks.append(_chunk(page, part, current))
                    part += 1
                    current = ""
                # Keep each Markdown table as its own evidence unit. This avoids
                # splitting row/column context or assigning a paragraph label to it.
                chunks.append(_chunk(page, part, paragraph))
                part += 1
            elif current and len(current) + len(paragraph) + 2 > target_size:
                chunks.append(_chunk(page, part, current))
                part += 1
                current = current[-overlap:] + "\n\n" + paragraph if overlap else paragraph
            else:
                current = f"{current}\n\n{paragraph}".strip()
        if current:
            chunks.append(_chunk(page, part, current))
    return chunks


def _chunk(page: ParsedPage, part: int, content: str) -> ParsedChunk:
    location = f"第 {page.page_number} 页，第 {part} 段" if page.extraction_method != "docx" else f"DOCX 第 {part} 段"
    content = clean_text(content)
    return ParsedChunk(
        page_start=page.page_number,
        page_end=page.page_number,
        location=location,
        structure_type="table" if content.startswith("|") else "paragraph",
        content=content,
        search_terms=tokenise_for_search(content),
        content_hash=hashlib.sha256(content.encode("utf-8")).hexdigest(),
    )


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
